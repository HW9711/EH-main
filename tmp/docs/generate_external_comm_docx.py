from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "doc" / "UART2外部通信控制协议对接说明.docx"

FRAME_HEAD = [0xD7, 0xCA, 0xF8, 0xF1]
FRAME_TAIL = [0xBF, 0xC6, 0xBC, 0xC4]
FRAME_FIXED_SIZE = 16
MAX_FRAME_SIZE = 150


def crc16_modbus_value(data: Sequence[int]) -> int:
    crc = 0xFFFF
    for byte in data:
        crc ^= byte & 0xFF
        for _ in range(8):
            if crc & 0x0001:
                crc = (crc >> 1) ^ 0xA001
            else:
                crc >>= 1
            crc &= 0xFFFF
    return crc


def be16(value: int) -> list[int]:
    return [(value >> 8) & 0xFF, value & 0xFF]


def build_frame(fun: int, area: int = 0xFF, info: int = 0xFF, payload: Sequence[int] | None = None, tran: int = 0x02) -> str:
    payload = list(payload or [])
    length = FRAME_FIXED_SIZE + len(payload)
    body = [tran, *be16(length), fun & 0xFF, area & 0xFF, info & 0xFF, *payload]
    crc = crc16_modbus_value(body)
    frame = [*FRAME_HEAD, *body, *be16(crc), *FRAME_TAIL]
    return " ".join(f"{b:02X}" for b in frame)


def set_cell_text(cell, text: str, bold: bool = False, code: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    if code:
        set_run_font(run, "Consolas", "Consolas", 9)
    else:
        set_run_font(run, "宋体", "Calibri", 9.5)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Calibri", size_pt: float | None = None) -> None:
    run.font.name = ascii_font
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 6, line: float = 1.15) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_para(doc: Document, text: str = "", style: str | None = None, bold: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(p)
    run = p.add_run(text)
    run.bold = bold
    set_run_font(run, "宋体", "Calibri", 10.5)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.2)
    set_paragraph_spacing(p, before=2, after=8, line=1.0)
    run = p.add_run(text)
    set_run_font(run, "Consolas", "Consolas", 9)
    run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    if level == 1:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        size = 16
    elif level == 2:
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        size = 13
    else:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        size = 11.5
    run.bold = True
    set_run_font(run, "微软雅黑", "Calibri", size)
    set_paragraph_spacing(p, before=10 if level == 1 else 6, after=4)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], code_cols: Iterable[int] = ()) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    code_cols = set(code_cols)
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], "1F3864")
        set_cell_text(header_cells[idx], header, bold=True)
        for paragraph in header_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, code=(idx in code_cols))
            if r_index % 2 == 1:
                set_cell_shading(cells[idx], "F7F9FC")
    doc.add_paragraph()


def add_bullets(doc: Document, items: Sequence[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, after=3)
        run = p.add_run(item)
        set_run_font(run, "宋体", "Calibri", 10.5)


def add_numbered(doc: Document, items: Sequence[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=3)
        run = p.add_run(item)
        set_run_font(run, "宋体", "Calibri", 10.5)


def add_toc(doc: Document) -> None:
    add_heading(doc, "目录", 1)
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run._r.append(instr)
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)
    placeholder = p.add_run("打开 Word 后右键更新目录，或按 F9 更新域。")
    set_run_font(placeholder, "宋体", "Calibri", 10.5)
    run_end = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)
    doc.add_page_break()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    for name in ["Normal", "Body Text"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("UART2外部通信控制协议对接说明 - ")
    set_run_font(r, "宋体", "Calibri", 9)
    r1 = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    r1._r.append(begin)
    r2 = footer.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r2._r.append(instr)
    r3 = footer.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r3._r.append(end)
    return doc


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    title = p.add_run("UART2外部通信控制协议对接说明")
    title.bold = True
    set_run_font(title, "微软雅黑", "Calibri", 24)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = p2.add_run("第三方设备外部控制联调版")
    subtitle.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    set_run_font(subtitle, "微软雅黑", "Calibri", 14)
    doc.add_paragraph()
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["适用接口", "USART2 外部通信控制链路"],
            ["文档目的", "第三方设备按本文组帧发送，即可申请外控、设置参数、执行启停、解析主控 ACK 和心跳。"],
            ["通信参数", "115200 bps，8 data bits，no parity，1 stop bit，no hardware flow control。"],
            ["协议范围", "V1 固定帧头/帧尾；Length 为整帧长度；CRC 覆盖 TranCode 到 InforArea。"],
            ["生成日期", "2026-06-08"],
        ],
    )
    add_para(doc, "重要说明：本文说明当前固件对外开放的 UART2 外部通信控制协议。第三方设备现场联调时，请按本文给出的串口参数、帧格式、CRC、外控保活和报警上传规则执行，并以实测收发结果确认对接状态。")
    doc.add_page_break()


def add_overview(doc: Document) -> None:
    add_heading(doc, "1. 对接目标和总体规则", 1)
    add_para(doc, "主控板支持第三方设备通过 UART2 发送协议帧实现外部通信控制。第三方设备只要满足串口电气连接、协议组帧、外控授权、保活和 ACK 解析规则，即可完成外部控制交互。")
    add_bullets(
        doc,
        [
            "所有下行帧必须使用 TranCode=0x02；主控上传和 ACK 使用 TranCode=0x01。",
            "所有设置、切换、普通启停命令必须先申请外部控制成功；急停是安全例外。",
            "固件收到合法下行帧后刷新链路保活计时；合法的含义是帧头、Length、帧尾、CRC 均通过。",
            "外控期间 1s 静默会先停电机和泵输出但保留外控权；5s 静默释放外控权。",
            "EEPROM 读写跟随当前选中 A/B 通道，第三方设备不能直接指定 I2C2/I2C3。",
            "心跳 InforArea 是动态长度，必须按在线字段逐段解析，不要按固定结构体硬拆。",
        ],
    )
    add_heading(doc, "1.1 第三方设备最小实现能力", 2)
    add_numbered(
        doc,
        [
            "能以 115200 8N1 在 USART2 物理链路收发原始字节。",
            "能按本文帧格式生成 Length、CRC16，并检查主控上传帧 CRC。",
            "能发送申请外控帧，并根据 0xDD/0xAA ACK 判断是否取得外控。",
            "外控成功后能定时发送保活帧，推荐 1s 周期；业务命令发送间隔小于 1s 时可省略额外保活。",
            "能按 ACK 码和失败原因处理 Busy、无通道、设备失败、长度错误等异常。",
            "能在异常、断链或联调结束时发送退出外控或急停帧。",
        ],
    )
    add_heading(doc, "1.2 对接边界说明", 2)
    add_bullets(
        doc,
        [
            "第三方设备只需要按 UART2 协议收发字节，不直接访问主控内部 I2C、EEPROM 总线或控制任务。",
            "主控上传给外设的是协议层状态和最终报警码；外设不能从报警帧反推出所有内部模块细节。",
            "普通控制命令必须在外控授权成功后发送；报警中、运行中或其它控制来源占用时，主控可能返回 BUSY。",
            "报警码为一个字节，外设应至少显示码值、中文含义和收到时间，方便现场追溯。",
        ],
    )


def add_serial_and_frame(doc: Document) -> None:
    add_heading(doc, "2. 串口和帧格式", 1)
    add_heading(doc, "2.1 串口参数", 2)
    add_table(
        doc,
        ["项目", "值", "说明"],
        [
            ["接口", "USART2", "主控外部通信专用口。"],
            ["TX/RX", "PD5 / PD6", "主控侧 USART2_TX/USART2_RX。第三方设备需交叉连接 RX/TX，并共地。"],
            ["波特率", "115200", "现场对接固定使用 115200 bps。"],
            ["数据位", "8", "UART_WORDLENGTH_8B。"],
            ["校验", "None", "UART_PARITY_NONE。"],
            ["停止位", "1", "UART_STOPBITS_1。"],
            ["硬件流控", "None", "UART_HWCONTROL_NONE。"],
            ["接收方式", "DMA 空闲包 + 软件 FIFO", "固件可处理半包、粘包、前导噪声和错误帧后重同步。"],
        ],
    )
    add_heading(doc, "2.2 固定帧结构", 2)
    add_code(doc, "Head(4) TranCode(1) Length(2) FunCode(1) AreaCode(1) InforCode(1) InforArea(N) CRC16(2) Tail(4)")
    add_table(
        doc,
        ["字段", "偏移", "长度", "取值/规则"],
        [
            ["Head", "0..3", "4", "固定 D7 CA F8 F1。解析器在 FIFO 内滑动查找。"],
            ["TranCode", "4", "1", "0x02=外部设备下发；0x01=主控上传/ACK。"],
            ["Length", "5..6", "2", "大端，表示整帧长度。最短 16，最大 150。"],
            ["FunCode", "7", "1", "命令大类。下行功能码见第 5 章。"],
            ["AreaCode", "8", "1", "同一 FunCode 下的动作、区域或页码。"],
            ["InforCode", "9", "1", "当前下行必须填 0xFF；上传帧按功能使用。"],
            ["InforArea", "10..CRC前", "N", "最大 134 字节。无载荷时 N=0。"],
            ["CRC16", "10+N..11+N", "2", "高字节在前；覆盖 TranCode 到 InforArea。"],
            ["Tail", "12+N..15+N", "4", "固定 BF C6 BC C4。"],
        ],
        code_cols=[1, 3],
    )
    add_para(doc, "第三方设备发送帧时，Length 必须等于 16 + InforArea 长度。不要把串口包长度、去掉帧头后的长度或仅 payload 长度写入 Length。")
    add_heading(doc, "2.3 下行帧通用模板", 2)
    add_code(doc, "D7 CA F8 F1 02 Length_H Length_L Fun Area FF [InforArea...] CRC_H CRC_L BF C6 BC C4")
    add_para(doc, "实际发送必须按大端：Length_H 在前，Length_L 在后。例如无载荷帧 Length=0x0010，应发送 00 10。")


def add_crc(doc: Document) -> None:
    add_heading(doc, "3. CRC16 计算方法", 1)
    add_para(doc, "本协议 CRC 与常见 CRC16/Modbus 位运算形式等价：初值 0xFFFF，多项式反射形式 0xA001，逐字节低位先移位。不同点是本协议帧内 CRC 按高字节、低字节发送。")
    add_code(
        doc,
        "uint16_t crc = 0xFFFF;\n"
        "for each byte b in [TranCode, Length_H, Length_L, FunCode, AreaCode, InforCode, InforArea...]:\n"
        "    crc ^= b;\n"
        "    repeat 8 times:\n"
        "        if (crc & 0x0001) crc = (crc >> 1) ^ 0xA001;\n"
        "        else              crc = crc >> 1;\n"
        "CRC_H = (crc >> 8) & 0xFF;\n"
        "CRC_L = crc & 0xFF;"
    )
    add_table(
        doc,
        ["校验项", "规则"],
        [
            ["覆盖范围", "从 TranCode 开始，到 InforArea 最后一个字节结束。帧头、CRC 自身、帧尾不参与。"],
            ["空 InforArea", "仍覆盖 TranCode、Length_H、Length_L、FunCode、AreaCode、InforCode 共 6 字节。"],
            ["字节序", "Length 和普通 16 位协议值大端；CRC 结果也按高字节、低字节放入帧。"],
            ["参考值", "ASCII `123456789` 的算法返回 0x4B37。"],
            ["调试建议", "先构造申请外控样例帧并让主控回 0xDD/0xAA，再联调其它命令。"],
        ],
    )


def add_state_machine(doc: Document) -> None:
    add_heading(doc, "4. 外部控制状态机", 1)
    add_heading(doc, "4.1 推荐联调顺序", 2)
    add_numbered(
        doc,
        [
            "连接串口，监听主控 100ms 心跳。能收到 `TranCode=0x01, FunCode=0xAA` 说明主控上传通道可用。",
            "发送申请外控 `FunCode=0x01, AreaCode=0xFF, InforCode=0xFF, InforArea=8字节授权码`。",
            "收到 `FunCode=0xDD, InforCode=0xAA` 后认为外控成功；收到 `0xAB` 或 `0xBB` 时不要发送启动类命令。",
            "按目标动作发送设置、切换、启停命令。设置速度/泵速只改设定，不等于启动输出。",
            "外控期间保持下行帧周期小于 1s。推荐每 1s 发送一次同一申请外控帧保活；业务命令本身也刷新固件保活计时。",
            "停止业务后发送 `FunCode=0xBB` 主动退出外控；异常时发送 `FunCode=0x04, AreaCode=0xFF` 急停。",
        ],
    )
    add_heading(doc, "4.2 外控授权和互斥", 2)
    add_table(
        doc,
        ["步骤/状态", "主控行为", "第三方设备动作"],
        [
            ["授权码长度不是 8", "返回 0xDD/0xBB，认证失败。", "检查 InforArea 长度。V1 只校验长度，不校验具体内容。"],
            ["脚踏/屏幕/手柄正在占用", "申请外控返回 0xDD/0xAB，失败原因 Busy。", "等待本机控制结束，或人工停止后重试。"],
            ["外控成功", "进入 CONTROL_OWNER_EXTERNAL，屏幕小电脑图标高亮。", "开始保活，允许设置和普通控制命令。"],
            ["未外控发送普通命令", "返回运行值/控制失败，原因 Busy。", "先重新申请外控。"],
            ["发送急停 0x04/0xFF", "允许跨来源强停并释放仲裁。", "用于安全停机，不要求已外控。"],
        ],
    )
    add_heading(doc, "4.3 两级链路看门狗", 2)
    add_table(
        doc,
        ["静默时间", "主控动作", "对接含义"],
        [
            [">=1000ms", "清电机运行、清 A/B 泵运行和速度，但保留外控授权。", "短暂停顿先安全停输出；链路恢复后仍可继续外控。"],
            [">=5000ms", "释放外控 owner，外控图标离线，本机控制可重新接管。", "第三方设备需重新申请外控。"],
            ["任意合法下行帧", "清零静默计时，刷新外部通信在线显示。", "保活帧、设置帧、启停帧都能续命。"],
        ],
    )
    add_para(doc, "注意：外控链路释放超时时间为 5000ms。第三方设备应把 1s 作为停输出保护边界，把 5s 作为外控权释放边界。")


def add_downlink_commands(doc: Document) -> None:
    add_heading(doc, "5. 下行命令详解", 1)
    add_heading(doc, "5.1 命令总表", 2)
    add_table(
        doc,
        ["FunCode", "名称", "AreaCode / InforArea", "是否要求外控", "成功响应"],
        [
            ["0x01", "申请外部控制", "Area=0xFF，Info=0xFF，InforArea=8字节授权码", "否", "0xDD / 0xAA"],
            ["0x02", "设置运行值", "速度、频率、A泵速度、B泵速度", "是", "0xDD / 0x01"],
            ["0x03", "切换设置", "A/B通道、方向、控制模式、工具类型", "是，且非运行/无报警", "0xDD / 0x01"],
            ["0x04", "控制命令", "泵启停、手柄启停、开口定位、急停", "急停否，其它是", "0xDD / 0x03"],
            ["0x05", "读业务 EEPROM 单页", "Area=0x01..0x08", "否，但需当前通道", "0x01 EEPROM上传帧"],
            ["0x06", "读业务 EEPROM 整区", "禁用", "否", "0xDD / 0x06 失败"],
            ["0x07", "写业务 EEPROM 单页", "Area=0x01..0x08，InforArea=30字节", "否，但需当前通道", "0xDD / 0x05"],
            ["0x08", "读导航 EEPROM 单页", "Area=Page12..128 或序号1..117", "否，但需当前通道", "0x01 EEPROM上传帧"],
            ["0x09", "读导航 EEPROM 整区", "禁用", "否", "0xDD / 0x06 失败"],
            ["0x0A", "写导航 EEPROM 单页", "Area=Page12..128 或序号1..117，InforArea=30字节", "否，但需当前通道", "0xDD / 0x05"],
            ["0xBB", "退出外控", "Area=0xFF，Info=0xFF，无载荷", "否", "0xDD / 0x03，载荷 BB"],
            ["0xFA", "权限开放", "Area=0xFF，Info=0xFF，InforArea=8字节权限码", "否", "0xDD / 0xFE"],
        ],
        code_cols=[0, 2, 4],
    )
    add_heading(doc, "5.2 设置运行值 FunCode=0x02", 2)
    add_table(
        doc,
        ["AreaCode", "含义", "InforArea", "成功 ACK 载荷", "失败条件"],
        [
            ["0x01", "当前手柄速度", "2字节大端。内部 `speed_work/speed_set_work` 使用 x10 单位。", "Area + 最终值2字节", "未外控、无当前通道、长度不足。"],
            ["0x02", "当前手柄往复频率", "支持 1字节 或 2字节大端。", "Area + 最终值2字节", "未外控、无当前通道、长度为0。"],
            ["0x03", "A泵速度", "2字节大端。只设置速度，不自动启动。", "Area + 最终值2字节", "未外控、长度不足。"],
            ["0x04", "B泵速度", "2字节大端。只设置速度，不自动启动。", "Area + 最终值2字节", "未外控、长度不足。"],
        ],
        code_cols=[0, 2, 3],
    )
    add_para(doc, "速度设置在运行中可动态生效；但启动手柄前应先下发非零速度，否则固件启动时会把 `speed_work` 恢复为 `speed_set_work`，可能仍为 0。")
    add_heading(doc, "5.3 切换设置 FunCode=0x03", 2)
    add_table(
        doc,
        ["AreaCode", "含义", "InforArea", "限制"],
        [
            ["0x01", "切换到 A 通道", "可空；ACK 回显 value=0。", "A 通道必须在线；运行中或报警中拒绝。"],
            ["0x02", "切换到 B 通道", "可空；ACK 回显 value=0。", "B 通道必须在线；运行中或报警中拒绝。"],
            ["0x03", "切换方向", "1字节：01=正转，02=反转，03=往复。", "必须非运行、无报警；方向值非法返回长度/参数错误。"],
            ["0x04", "切换控制模式", "1字节：01=脚踏，02=手控，03=外部控制。", "只接受这三类值。"],
            ["0x05", "切换工具类型", "1字节：01=刨头，02=磨头。", "只接受这两类值。"],
        ],
        code_cols=[0, 2],
    )
    add_heading(doc, "5.4 控制命令 FunCode=0x04", 2)
    add_table(
        doc,
        ["AreaCode", "动作", "InforArea", "主控实际行为", "失败条件"],
        [
            ["0x01", "A泵启动", "空", "A泵在线时置上位机独立 A 泵运行请求；若速度为0，按泵类型补默认速度。", "未外控、A泵离线。"],
            ["0x02", "A泵停止", "空", "清上位机独立 A 泵请求；若手柄冷却跟随仍选中 A，A 泵可能继续运行。", "未外控。"],
            ["0x03", "B泵启动", "空", "B泵在线时置运行，速度为0时补默认速度。", "未外控、B泵离线。"],
            ["0x04", "B泵停止", "空", "清 B泵运行和排空计时。", "未外控。"],
            ["0x05", "当前手柄启动", "空", "要求当前通道在线；把 `speed_work` 恢复为设定速度；可联动注水泵冷却。", "未外控、无在线手柄、报警。"],
            ["0x06", "当前手柄停止", "空", "清手柄运行、速度和外控启动标志；停止手柄带动的注水泵跟随。", "未外控。"],
            ["0x07", "开口定位左", "空", "按当前选中通道执行 ToolPosMay 左动作。", "无当前通道。"],
            ["0x08", "开口定位右", "空", "按当前选中通道执行 ToolPosMay 右动作。", "无当前通道。"],
            ["0xFF", "急停/紧急刹车", "空", "强制停电机、A/B泵，清外控/触控/脚踏控制标志，并释放仲裁。", "安全例外，允许未外控发送。"],
        ],
        code_cols=[0, 2],
    )
    add_para(doc, "A泵停止和手柄冷却跟随有合并逻辑：如果 A 泵是当前规则选中的注水冷却泵，即使发送 A泵停止，手柄运行期间 A 泵仍可能由冷却跟随保持运行。需要停止手柄或关闭跟随来源。")
    add_heading(doc, "5.5 EEPROM 读写命令", 2)
    add_table(
        doc,
        ["AreaCode", "业务区", "EEPROM Page", "读写说明"],
        [
            ["0x01", "识别信息区", "Page2", "读返回前30字节；写要求30字节。"],
            ["0x02", "刀具信息区", "Page3", "读返回前30字节；写要求30字节。"],
            ["0x03", "初始值信息区", "Page4", "常用于默认速度/频率/泵流量等。"],
            ["0x04", "按键自定义区", "Page5", "读写当前通道 EEPROM。"],
            ["0x05", "多档位调节区", "Page6", "Page7 保留，不暴露。"],
            ["0x06", "运行信息储存区", "Page8", "读写当前通道 EEPROM。"],
            ["0x07", "外部编辑区", "Page9", "Page10 保留，不暴露。"],
            ["0x08", "出厂信息区", "Page11", "读写当前通道 EEPROM。"],
        ],
        code_cols=[0, 2],
    )
    add_bullets(
        doc,
        [
            "读业务页：下发 FunCode=0x05，AreaCode=上表区域，主控上传 FunCode=0x01、AreaCode 回显、InforCode=0xFF、InforArea=30字节。",
            "写业务页：下发 FunCode=0x07，InforArea 必须刚好 30字节；底层驱动自动生成页尾 2字节校验。",
            "读导航页：下发 FunCode=0x08，AreaCode 可直接用 Page12..128，也可用导航序号 1..117；上传 AreaCode=0xFF，InforCode 回显页码。",
            "写导航页：下发 FunCode=0x0A，规则同读导航页，InforArea 必须 30字节。",
            "整区读取 0x06 和 0x09 当前 V1 禁用，会返回 NOT_SUPPORT。",
            "EEPROM 操作永远使用当前选中通道。批量写入前必须先切到目标 A/B 通道并确认心跳选中字段。",
        ],
    )


def add_upload_and_ack(doc: Document) -> None:
    add_heading(doc, "6. 主控上传帧和 ACK 解析", 1)
    add_heading(doc, "6.1 ACK 帧 FunCode=0xDD", 2)
    add_para(doc, "ACK 帧统一使用上传方向：TranCode=0x01，FunCode=0xDD，AreaCode=0xFF，InforCode 为 ACK 码，InforArea 为可选回显或失败原因。")
    add_table(
        doc,
        ["InforCode", "含义", "常见 InforArea"],
        [
            ["0x01", "运行值设置成功", "AreaCode + 2字节最终值。"],
            ["0x02", "运行值设置失败", "失败对象 + 失败原因。"],
            ["0x03", "控制命令成功", "控制 AreaCode；退出外控时为 BB。"],
            ["0x04", "控制命令失败", "失败对象 + 失败原因。"],
            ["0x05", "EEPROM 读写成功", "写入成功时回显 AreaCode。"],
            ["0x06", "EEPROM 读写失败", "失败对象 + 失败原因。"],
            ["0xAA", "外部控制申请成功", "空。"],
            ["0xAB", "外部控制申请失败", "通常为 01 06，表示申请命令 Busy。"],
            ["0xBB", "授权码校验失败", "空。当前 V1 多为长度不是8。"],
            ["0xFE", "权限开放成功", "空。"],
            ["0xFF", "权限开放失败", "空。当前 V1 多为长度不是8。"],
        ],
        code_cols=[0, 2],
    )
    add_heading(doc, "6.2 失败原因码", 2)
    add_table(
        doc,
        ["原因码", "名称", "说明", "对接处理"],
        [
            ["0x01", "BAD_LENGTH", "InforArea 长度不符合命令要求。", "检查载荷字节数。"],
            ["0x02", "BAD_AREA", "AreaCode、FunCode 或取值非法。", "检查命令矩阵和枚举值。"],
            ["0x03", "NO_CHANNEL", "无当前选中 A/B，或目标通道不在线。", "等待手柄识别，或先切换到在线通道。"],
            ["0x04", "DEVICE_FAIL", "EEPROM/I2C 读写失败、泵未在线或设备未识别。", "检查硬件在线和当前通道。"],
            ["0x05", "NOT_SUPPORT", "当前 V1 明确不支持。", "不要使用整区读取等大包命令。"],
            ["0x06", "BUSY", "未取得外控、运行中、报警中或其它来源占用。", "停止本机来源、清报警或重新申请外控。"],
        ],
        code_cols=[0],
    )
    add_heading(doc, "6.3 报警上传 FunCode=0x03", 2)
    add_para(doc, "主控会在报警值变化时上传主机运行内容：TranCode=0x01，FunCode=0x03，AreaCode=0xFF，InforCode=0x05，InforArea[0] 为报警码。0 表示无报警或报警解除；非 0 表示当前需要外设提示的报警。第三方设备应把该帧作为弹窗、日志或状态告警依据。")
    add_para(doc, "报警上传不是周期心跳字段，而是变化触发同步：同一个报警值保持不变时不会每个周期重复刷帧。运行中另一路手柄校验失败这类临时报警可能只保持约 3 秒，到期后主控会再上传 0x00 关闭提示。")
    add_table(
        doc,
        ["报警码", "十进制", "对接显示建议", "来源和处理说明"],
        [
            ["0x00", "0", "无报警 / 报警解除", "外设收到 0x00 后应关闭当前报警弹窗或把状态标记为恢复。"],
            ["0x01", "1", "手柄未连接，请连接手柄", "当前需要手柄参与的操作没有检测到有效手柄。"],
            ["0x02", "2", "手控已选中，请用手控", "本机手控来源占用，外部设备普通控制应等待或先退出当前来源。"],
            ["0x03", "3", "脚控已选中，请用脚控 / 电机相位错误", "普通控制路径表示脚控占用；若由电机驱动 Err=14 映射而来，则表示缺相/相位错误。报警帧本身不带来源字段，外设应记录码值并结合心跳运行状态判断。"],
            ["0x04", "4", "电机过载 / 电机霍尔错误", "普通控制路径表示电机过载；若由驱动 Err=11/12 映射而来，则表示霍尔断线或霍尔学习错误。"],
            ["0x05", "5", "脚踏值错误 / 电机过载或刀具卡住", "普通控制路径表示脚踏值错误；若由驱动 Err=2/5 映射而来，则表示过流或运行中堵转。"],
            ["0x06", "6", "电机过载，请松开脚踏", "含义按电机过载处理，外设提示用户停止当前动作并检查负载。"],
            ["0x07", "7", "UID 错误", "RFID/UID 识别相关异常。"],
            ["0x08", "8", "电机通讯异常 / 系统供电电压异常", "普通报警表示电机通信异常；若由驱动 Err=3/4 映射而来，则表示过压或欠压。"],
            ["0x09", "9", "HALL 值错误", "手柄或电机相关 HALL 值异常。"],
            ["0x0A", "10", "A 通道手柄型号错误", "A 通道 EEPROM 校验失败或型号校验失败。"],
            ["0x0B", "11", "驱动板故障，请联系售后", "外部通信收到 0x0B 时按驱动板故障处理；Page4 速度/频率阈值也使用 0x0B 触发蜂鸣，但该阈值报警通常不写入外部通信报警帧。"],
            ["0x0C", "12", "B 通道手柄型号错误", "B 通道 EEPROM 校验失败或型号校验失败。"],
            ["0x0D", "13", "运行中手柄插拔/掉线报警", "运行过程中检测到手柄基座拔出或掉线，外设应提示用户停止操作并检查手柄连接。"],
            ["0x0E", "14", "A/B 通道手柄型号均错误", "A、B 两路都存在 EEPROM 校验失败或型号校验失败。"],
        ],
        code_cols=[0],
    )
    add_para(doc, "电机驱动板原始 Err 不是 UART2 报警上传帧中的字段。主控会先把驱动 Err 转换成上表中的一个字节报警码，再通过 InforArea[0] 上传；因此第三方设备不能只凭报警帧恢复出原始 Err。")
    add_table(
        doc,
        ["驱动 Err", "驱动含义", "UART2 最终报警码", "外设显示建议"],
        [
            ["0", "无故障", "0x00", "清除本次驱动报警。"],
            ["1", "模块保护", "0x0B", "驱动板故障。"],
            ["2", "过流保护", "0x05", "电机过载或刀具卡住。"],
            ["3", "过压保护", "0x08", "系统供电电压异常。"],
            ["4", "欠压保护", "0x08", "系统供电电压异常。"],
            ["5", "运行中堵转", "0x05", "电机过载或刀具卡住。"],
            ["6", "驱动器过温", "0x0B", "驱动板故障。"],
            ["7", "参数保存错误", "0x0B", "驱动板故障。"],
            ["8", "刹车时间过长", "0x0B", "驱动板故障。"],
            ["9", "编码器错误", "0x0B", "驱动板故障。"],
            ["10", "开环检测错误", "0x0B", "驱动板故障。"],
            ["11", "霍尔断线", "0x04", "电机霍尔错误。"],
            ["12", "霍尔学习错误", "0x04", "电机霍尔错误。"],
            ["13", "通信握手错误", "0x0B", "驱动板故障。"],
            ["14", "缺相", "0x03", "电机相位错误。"],
            ["15", "Hall 拖动错误", "0x0B", "驱动板故障。"],
        ],
        code_cols=[0, 2],
    )


def add_heartbeat(doc: Document) -> None:
    add_heading(doc, "7. 心跳帧 FunCode=0xAA", 1)
    add_para(doc, "主控每 100ms 主动上传一次心跳。心跳用于第三方设备显示 A/B 手柄、当前选中通道、运行状态、脚踏、A/B 泵、压力扩展等状态。心跳不是第三方设备喂给主控的保活；主控外控看门狗只看合法下行帧。")
    add_heading(doc, "7.1 心跳基本帧", 2)
    add_code(doc, "TranCode=0x01, FunCode=0xAA, AreaCode=0xFF, InforCode=0xFF, InforArea=动态长度")
    add_heading(doc, "7.2 动态字段解析顺序", 2)
    add_table(
        doc,
        ["顺序", "字段", "是否固定出现", "解析规则"],
        [
            ["1", "A手柄在线状态", "固定", "0x01 在线，0xFF 离线。"],
            ["2", "A手柄原始类型2字节", "A在线时出现", "来自 EEPROM Page2 前两字节，例如 6B 01。"],
            ["3", "B手柄在线状态", "固定", "0x01 在线，0xFF 离线。"],
            ["4", "B手柄原始类型2字节", "B在线时出现", "来自 EEPROM Page2 前两字节。"],
            ["5", "当前选中手柄", "固定", "0x01=A，0x02=B，0xFF=未选中或离线。"],
            ["6", "当前手柄运行状态", "固定", "0x01=待机，0x02=运行中，0x03=未接入。"],
            ["7", "当前手柄模式", "固定", "0x01=正转，0x02=反转，0x03=往复，0xFF=未知。"],
            ["8", "当前速度2字节大端", "仅运行中出现", "`WorkMessage.speed_work`，内部 x10；显示 rpm 时一般除以10。"],
            ["9", "驱动电流2字节大端", "仅运行中出现", "`driver_current_x100`，单位 0.01A。"],
            ["10", "脚踏在线状态", "固定", "0x01 在线，0xFF 离线。"],
            ["11", "A泵在线状态", "固定", "0x01 在线，0xFF 离线。"],
            ["12", "A泵类型+速度+压力扩展", "A泵在线时出现", "类型1字节，速度2字节大端，压力扩展11字节。"],
            ["13", "B泵在线状态", "固定", "0x01 在线，0xFF 离线。"],
            ["14", "B泵类型+速度+压力扩展", "B泵在线时出现", "类型1字节，速度2字节大端，压力扩展11字节。"],
        ],
        code_cols=[0, 3],
    )
    add_heading(doc, "7.3 压力扩展 11 字节", 2)
    add_para(doc, "泵在线时，心跳中每路泵在类型和速度后追加 11 字节压力扩展。该扩展复用 CS1237 下位机字段顺序，内部为 little-endian，不按外部通信大端规则。")
    add_table(
        doc,
        ["偏移", "字段", "长度", "字节序", "说明"],
        [
            ["0..3", "RawCs1237", "4", "little-endian", "int32 原始采样二进制位。"],
            ["4..7", "WeightX10", "4", "little-endian", "重量，单位 0.1g。"],
            ["8..9", "ThresholdG", "2", "little-endian", "压力阈值，单位 g。"],
            ["10", "Seq", "1", "单字节", "压力模块帧序号。"],
        ],
        code_cols=[0, 3],
    )
    add_para(doc, "解析建议：使用游标按在线状态逐段推进。若 A 或 B 离线，不要读取该路泵的类型、速度、压力扩展，否则后续字段会错位。")


def add_examples(doc: Document) -> None:
    add_heading(doc, "8. 完整十六进制样例帧", 1)
    add_para(doc, "以下样例按默认授权码 `11 22 33 44 55 66 77 88` 生成，CRC 已按本文规则复算。第三方设备可先逐条发送这些帧验证链路。")
    examples = [
        ["申请外部控制", "先发；成功 ACK 为 0xDD/0xAA", build_frame(0x01, payload=[0x11, 0x22, 0x33, 0x44, 0x55, 0x66, 0x77, 0x88])],
        ["权限开放", "可选；成功 ACK 为 0xDD/0xFE", build_frame(0xFA, payload=[0x11, 0x22, 0x33, 0x44, 0x55, 0x66, 0x77, 0x88])],
        ["设置当前手柄速度 3000", "3000=0x0BB8，大端", build_frame(0x02, 0x01, payload=[0x0B, 0xB8])],
        ["设置频率 20", "1字节频率", build_frame(0x02, 0x02, payload=[0x14])],
        ["切换到 A 通道", "无载荷", build_frame(0x03, 0x01)],
        ["切换到 B 通道", "无载荷", build_frame(0x03, 0x02)],
        ["方向正转", "value=0x01", build_frame(0x03, 0x03, payload=[0x01])],
        ["方向反转", "value=0x02", build_frame(0x03, 0x03, payload=[0x02])],
        ["方向往复", "value=0x03", build_frame(0x03, 0x03, payload=[0x03])],
        ["模式外部控制", "value=0x03", build_frame(0x03, 0x04, payload=[0x03])],
        ["工具类型刨头", "value=0x01", build_frame(0x03, 0x05, payload=[0x01])],
        ["A 泵启动", "要求 A 泵在线且已外控", build_frame(0x04, 0x01)],
        ["A 泵停止", "清独立 A 泵请求", build_frame(0x04, 0x02)],
        ["B 泵启动", "要求 B 泵在线且已外控", build_frame(0x04, 0x03)],
        ["B 泵停止", "清 B 泵运行", build_frame(0x04, 0x04)],
        ["当前手柄启动", "要求当前手柄在线、无报警、已外控", build_frame(0x04, 0x05)],
        ["当前手柄停止", "停止电机和手柄冷却跟随", build_frame(0x04, 0x06)],
        ["开口定位左", "要求当前通道有效", build_frame(0x04, 0x07)],
        ["开口定位右", "要求当前通道有效", build_frame(0x04, 0x08)],
        ["急停", "安全例外，未外控也允许", build_frame(0x04, 0xFF)],
        ["退出外控", "释放外控权", build_frame(0xBB)],
        ["读取业务 Page4", "Fun=0x05, Area=0x03", build_frame(0x05, 0x03)],
        ["读取导航 Page12", "Fun=0x08, Area=0x0C", build_frame(0x08, 0x0C)],
    ]
    add_table(doc, ["动作", "说明", "完整帧"], examples, code_cols=[2])
    add_heading(doc, "8.1 推荐动作序列", 2)
    add_code(
        doc,
        "1. 申请外控 -> 等待 0xDD/0xAA\n"
        "2. 切换到目标 A/B 通道 -> 等待 0xDD/0x01\n"
        "3. 设置速度/频率/泵速度 -> 等待 0xDD/0x01\n"
        "4. 发送当前手柄启动或泵启动 -> 等待 0xDD/0x03\n"
        "5. 每 <1s 发送业务帧或申请外控帧保活\n"
        "6. 停止手柄/泵 -> 等待 0xDD/0x03\n"
        "7. 退出外控 -> 等待 0xDD/0x03, InforArea=BB"
    )
    add_heading(doc, "8.2 写 EEPROM 页模板", 2)
    add_para(doc, "写业务页和写导航页的 InforArea 必须是 30 字节。下面只展示结构，实际 30 字节内容按 EEPROM 布局填写。CRC 需用完整 30 字节重新计算。")
    add_code(doc, "业务页写入：D7 CA F8 F1 02 00 2E 07 [Area] FF [30字节数据] CRC_H CRC_L BF C6 BC C4\n导航页写入：D7 CA F8 F1 02 00 2E 0A [Page或序号] FF [30字节数据] CRC_H CRC_L BF C6 BC C4")


def add_testing(doc: Document) -> None:
    add_heading(doc, "9. 联调验收和故障定位", 1)
    add_heading(doc, "9.1 最小验收用例", 2)
    add_table(
        doc,
        ["用例", "操作", "预期"],
        [
            ["串口连通", "仅监听主控上传", "100ms 左右收到 0xAA 心跳。"],
            ["CRC 正确性", "发送申请外控样例帧", "收到 0xDD/0xAA 或 0xDD/0xAB；若完全无 ACK，先查 CRC/帧尾/波特率。"],
            ["外控互斥", "本机脚踏或手柄占用时申请外控", "返回 0xDD/0xAB，失败原因 Busy。"],
            ["设置速度", "外控成功后发送设置速度 3000", "返回 0xDD/0x01，InforArea=01 0B B8。"],
            ["启动手柄", "当前通道在线、已设速度后发送 0x04/0x05", "返回 0xDD/0x03，心跳运行状态变 0x02。"],
            ["断链短超时", "外控运行时停止所有下行 >1s", "电机/泵停输出，但外控图标和授权短期保留。"],
            ["断链长超时", "继续静默 >5s", "释放外控，第三方设备需重新申请。"],
            ["急停", "未申请外控也发送 0x04/0xFF", "主控强停并 ACK 0xDD/0x03。"],
            ["EEPROM读页", "切目标通道后读 Page4", "收到 0x01 上传帧，InforArea 为 30 字节。"],
        ],
    )
    add_heading(doc, "9.2 常见故障", 2)
    add_table(
        doc,
        ["现象", "优先检查", "说明"],
        [
            ["主控无任何响应", "波特率是否 115200、TX/RX 是否交叉、CRC 是否高字节在前", "旧资料 9600 可能误导。"],
            ["有心跳但申请外控无 ACK", "下行 TranCode 是否 0x02，InforCode 是否 0xFF，Length 是否整帧长度", "上传和下行方向码不能混用。"],
            ["收到 0xBB 授权失败", "授权码长度是否 8 字节", "V1 只校验长度。"],
            ["收到 Busy", "是否未外控、运行中、报警中或本机来源占用", "先看 ACK 失败原因第 2 字节。"],
            ["设速成功但手柄不转", "是否已发送启动命令、当前通道是否在线、速度是否非零", "设置不等于启动。"],
            ["A泵停止后仍转", "是否手柄冷却跟随仍选中 A 注水泵", "停止手柄或清跟随来源。"],
            ["心跳字段错位", "是否按在线字段动态解析", "离线设备不会追加类型、速度、压力扩展。"],
            ["EEPROM读写失败", "当前是否选中通道、手柄 EEPROM 是否在线、Area 是否合法", "主控不能由外设直接指定 A/B EEPROM 总线。"],
            ["运行中切方向失败", "运行中/报警中禁止切换", "先停止手柄，再切换。"],
        ],
    )
    add_heading(doc, "9.3 第三方设备实现建议", 2)
    add_bullets(
        doc,
        [
            "发送侧维护一个统一组帧函数，所有命令都由该函数填 Length 和 CRC，避免每个按钮单独拼 CRC。",
            "接收侧先滑动查找帧头，再读取 Length 判断完整帧，再校验帧尾和 CRC；不要假定一次串口 read 就是一帧。",
            "外控授权状态以 0xDD/0xAA ACK 为准；心跳只表示主控在线，不会替第三方设备喂固件下行看门狗。",
            "保活帧推荐复用申请外控帧，并在成功申请后锁定授权码，避免用户修改输入导致后续保活长度或 CRC 变化。",
            "控制命令和保活共用下行链路时，若刚发送过业务命令，可跳过本周期保活，降低串口拥堵。",
            "所有启动类动作前先确认心跳中的当前通道、运行状态、报警码和设备在线状态。",
            "EEPROM 写入要有二次确认和写后读回；运行中不建议批量写 EEPROM。",
        ],
    )


def add_appendix(doc: Document) -> None:
    add_heading(doc, "10. 字段速查", 1)
    add_table(
        doc,
        ["类别", "值", "含义"],
        [
            ["TranCode", "0x01", "主控上传或 ACK。"],
            ["TranCode", "0x02", "外部设备下发。"],
            ["通用空值", "0xFF", "无区域码、无信息码、离线或未知。"],
            ["在线状态", "0x01", "在线。"],
            ["在线状态", "0xFF", "离线。"],
            ["选中通道", "0x01 / 0x02 / 0xFF", "A / B / 未选中。"],
            ["运行状态", "0x01 / 0x02 / 0x03", "待机 / 运行中 / 未接入。"],
            ["手柄模式", "0x01 / 0x02 / 0x03 / 0xFF", "正转 / 反转 / 往复 / 未知。"],
            ["控制模式", "0x01 / 0x02 / 0x03", "脚踏 / 手控 / 外部控制。"],
            ["工具类型", "0x01 / 0x02", "刨头 / 磨头。"],
        ],
        code_cols=[1],
    )


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    add_cover(doc)
    add_toc(doc)
    add_overview(doc)
    add_serial_and_frame(doc)
    add_crc(doc)
    add_state_machine(doc)
    add_downlink_commands(doc)
    add_upload_and_ack(doc)
    add_heartbeat(doc)
    add_examples(doc)
    add_testing(doc)
    add_appendix(doc)

    section = doc.sections[-1]
    section.start_type = WD_SECTION.NEW_PAGE
    doc.core_properties.title = "UART2外部通信控制协议对接说明"
    doc.core_properties.author = "Codex"
    doc.core_properties.subject = "第三方设备外部控制联调"
    doc.core_properties.comments = "Generated for third-party UART2 external-control integration on 2026-06-08."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
